"""Deterministic seeds and independently readable synthetic fixtures.

Optional generators are isolated: absence/errors become manifest gaps, never fake
positive fixtures. ZIP aliases are explicitly container-only, not app validity.
"""
import ast
import bz2
import gzip
import io
import json
import lzma
import random
import shutil
import sqlite3
import struct
import tempfile
import zlib
from pathlib import Path

from .common import file_inventory, write_json
from .oracles import command
from .special_formats import car_profile, cpio_bz2, cpio_profile, crx3

FIXED_TIME = (2020, 1, 2, 3, 4, 6)
JSON_EXTS = 'json geojson ipynb map har topojson gltf'.split()
XML_EXTS = 'xml ui rels xhtml kml gpx dae rss atom xmp xsl xslt fb2'.split()
STREAMS = 'gz xz bz2 zst br lz4 lz lzma lzo z'.split()


def zip_bytes(entries, level=0, mimetype=None):
    buffer = io.BytesIO()
    with zipfile_factory(buffer, level) as archive:
        for name, data in entries.items():
            import zipfile
            info = zipfile.ZipInfo(name, FIXED_TIME)
            info.external_attr = (0o40755 if name.endswith('/') else 0o100644) << 16
            info.compress_type = zipfile.ZIP_STORED if name == mimetype or level == 0 else zipfile.ZIP_DEFLATED
            archive.writestr(info, data, compresslevel=level or None)
    return buffer.getvalue()


def zipfile_factory(buffer, level):
    import zipfile
    return zipfile.ZipFile(buffer, 'w', compression=zipfile.ZIP_DEFLATED if level else zipfile.ZIP_STORED)


def compress(data, codec, level=1):
    if codec == 'gz':
        return gzip.compress(data, compresslevel=level, mtime=0)
    if codec == 'bz2':
        return bz2.compress(data, compresslevel=level)
    if codec in ('xz', 'lzma'):
        return lzma.compress(data, preset=level, format=lzma.FORMAT_XZ if codec == 'xz' else lzma.FORMAT_ALONE)
    if codec == 'zst':
        import zstandard
        return zstandard.ZstdCompressor(level=level).compress(data)
    if codec == 'br':
        import brotli
        return brotli.compress(data, quality=level)
    if codec == 'lz4':
        import lz4.frame
        return lz4.frame.compress(data, compression_level=level)
    if codec == 'z':
        return command(['compress', '-c'], input=data)
    exe = {'lz': 'lzip', 'lzo': 'lzop'}[codec]
    return command([exe, '-c', '-1'], input=data)


class Builder:
    def __init__(self, root, catalog, scale=1):
        self.root = Path(root)
        self.directory = self.root / 'corpus/generated'
        self.directory.mkdir(parents=True, exist_ok=True)
        self.cases = []
        self.gaps = []
        self.catalog = catalog
        self.scale = scale
        rng = random.Random(1729)
        self.noise = rng.randbytes(32768 * scale)
        self.payload = b''.join((f'{i:06d}: sector={i % 23}, value=synthetic, label=benchmark\n'.encode()
                                 for i in range(4096 * scale)))
        self.json = json.dumps({'label': 'Synthetic corpus; no personal data',
                                'records': [{'id': i, 'group': i % 11, 'label': 'repeat label'}
                                            for i in range(256 * scale)]}, indent=4).encode()
        self.xml = (b'<?xml version="1.0" encoding="UTF-8"?>\n<!-- preserve comment -->\n'
                    b'<data  version = "1"><label> keep  spaces </label>' +
                    b'<record  id = "1"  label = "repeat" />' * (256 * scale) + b'</data>')
        self.children = {'data.json': {'kind': 'json'}, 'data.xml': {'kind': 'xml'},
                         'nested/records.jsonl': {'kind': 'jsonl'}}
        self.entries = {'data.json': self.json, 'data.xml': self.xml,
                        'nested/records.jsonl': b' { "n": 1, "text": "keep  spaces" } \r\n' * 100,
                        'payload.bin': self.payload, 'noise.bin': self.noise,
                        'empty/': b'', 'zero.bin': b''}

    def add(self, name, data, ext, family, oracle, *, scope='native', tier='extended',
            expectation='observe', provenance=None, license='BSD-3-Clause', note=None):
        path = self.directory / name
        if data is not None:
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(data)
        case = {'id': name.replace('.', '-').replace('/', '-').lower(),
                'path': str(path.relative_to(self.root)), 'extension': ext, 'family': family,
                'scope': scope, 'tier': tier, 'oracle': oracle, 'expectation': expectation,
                'license': license, 'provenance': provenance or
                {'type': 'generated', 'generator': 'frbench.generate', 'seed': 1729,
                 'description': note or 'Synthetic data authored for this corpus'},
                'inventory': file_inventory(path)}
        self.cases.append(case)
        return path

    def optional(self, name, function):
        mark = len(self.cases)
        try:
            function()
        except Exception as exc:
            # A family is atomic: partial files are not called complete coverage.
            self.cases = self.cases[:mark]
            self.gaps.append({'generator': name, 'reason': type(exc).__name__ + ': ' + str(exc)})
            print('GAP', name, self.gaps[-1]['reason'], flush=True)

    def text(self):
        for ext in JSON_EXTS:
            payload = self.json
            if ext == 'geojson':
                payload = json.dumps({'type': 'FeatureCollection', 'features': [
                    {'type': 'Feature', 'properties': {'name': 'Synthetic'},
                     'geometry': {'type': 'Point', 'coordinates': [30, 50]}}]}, indent=4).encode()
            elif ext == 'ipynb':
                payload = json.dumps({'nbformat': 4, 'nbformat_minor': 5, 'metadata': {},
                                      'cells': [{'cell_type': 'markdown', 'id': 'synthetic',
                                                 'metadata': {}, 'source': ['# Synthetic\n']}]}, indent=4).encode()
            elif ext == 'gltf':
                payload = json.dumps({'asset': {'version': '2.0', 'generator': 'frbench'},
                                      'scenes': [{'nodes': []}], 'scene': 0}, indent=4).encode()
            elif ext == 'map':
                payload = b' { "version": 3, "sources": [], "names": [], "mappings": "" } '
            elif ext == 'har':
                payload = json.dumps({'log': {'version': '1.2', 'creator': {'name': 'frbench',
                                                                          'version': '1'},
                                             'entries': []}}, indent=4).encode()
            elif ext == 'topojson':
                payload = b' { "type": "Topology", "objects": {}, "arcs": [] } '
            self.add('text/pretty.' + ext, payload, ext, 'text', {'kind': 'json'}, tier='smoke' if ext == 'json' else 'extended')
        self.add('text/exact-tokens.json', b' { "n": 1.2300E+99999, "n": -0, "s": "keep  \\u0061" } ',
                 'json', 'text', {'kind': 'json'})
        for ext in ['jsonl', 'ndjson']:
            self.add('text/records.' + ext, self.entries['nested/records.jsonl'], ext,
                     'text', {'kind': 'jsonl'})
        for ext in XML_EXTS:
            # Syntax profiles are explicit: a generic XML tree is not a COLLADA/FB2 validator.
            self.add('text/attributes.' + ext, self.xml, ext, 'text', {'kind': 'xml'},
                     scope='syntax', tier='smoke' if ext == 'xml' else 'extended')
        qgs = (b'<!DOCTYPE qgis PUBLIC \'http://mrcc.com/qgis.dtd\' \'SYSTEM\'>\n'
               b'<qgis  version = "3.44.0" projectname = "Synthetic">'
               b'<title> keep  spaces </title><projectlayers/></qgis>')
        self.add('text/project.qgs', qgs, 'qgs', 'text', {'kind': 'xml'})
        self.add('text/compact.json', json.dumps(json.loads(self.json), separators=(',', ':')).encode(),
                 'json', 'text', {'kind': 'json'}, note='Already minified control')
        for ext, data in [('json', b'{broken'), ('jsonl', b'1\n\n2\n'), ('xml', b'<unclosed'),
                          ('map', b'linker map, not a JSON source map')]:
            self.add('controls/malformed.' + ext, data, ext, 'text', {'kind': 'bytes'},
                     tier='control', expectation='unchanged')
        return qgs

    def streams(self):
        for codec in STREAMS:
            def create(codec=codec):
                for name, data, level in [('weak', self.payload, 1), ('entropy', self.noise, 9),
                                          ('compact', self.payload, 9)]:
                    self.add(f'streams/{name}.{codec}', compress(data, codec, level), codec,
                             'streams', {'kind': 'stream', 'codec': codec},
                             tier='smoke' if codec == 'gz' and name == 'weak' else 'extended')
            self.optional(codec, create)

    def archives(self, qgs):
        import tarfile
        from .catalog import SPECIAL_FAMILY
        for ext in self.catalog['archive_extensions']:
            family = SPECIAL_FAMILY.get(ext, 'zip')
            if family != 'zip':
                continue
            entries = dict(self.entries)
            oracle = {'kind': 'zip', 'children': self.children}
            scope = 'container-only'
            if ext == 'qgz':
                entries = {'project.qgs': qgs, 'readme.txt': b'Synthetic QGIS project'}
                oracle = {'kind': 'zip', 'children': {'project.qgs': {'kind': 'xml'}}}
                scope = 'native'
            elif ext == 'mellel':
                entries['main.xml'] = b'<archive creator="com.redlex.mellel" writer-version="1" compatibility-version="1"><text> keep text </text></archive>'
                oracle['children'] = {}
                scope = 'syntax'
            elif ext in ('usdz', 'crx'):
                # Real USDZ requires 64-byte alignment/STORED; CRX requires a signed header.
                continue
            elif ext == 'zip':
                scope = 'native'
            self.add('archives/container.' + ext, zip_bytes(entries), ext, 'archives', oracle,
                     scope=scope, tier='smoke' if ext == 'zip' else 'extended',
                     note='Container transport fixture; no application validity claim' if scope == 'container-only' else None)
        self.add('archives/weak.zip', zip_bytes(self.entries, level=1), 'zip', 'archives',
                 {'kind': 'zip', 'children': self.children})
        buffer = io.BytesIO()
        with tarfile.open(fileobj=buffer, mode='w', format=tarfile.PAX_FORMAT) as archive:
            for name, data in self.entries.items():
                info = tarfile.TarInfo(name)
                info.mtime, info.mode, info.uid, info.gid = 1577934246, 0o644, 1000, 1000
                info.uname = info.gname = 'synthetic'
                info.size = len(data)
                if name.endswith('/'):
                    info.type, info.mode = tarfile.DIRTYPE, 0o755
                archive.addfile(info, io.BytesIO(data))
        tar = buffer.getvalue()
        for ext in ['tar', 'cbt', 'gem']:
            self.add('archives/container.' + ext, tar, ext, 'archives',
                     {'kind': 'tar', 'children': self.children}, scope='native' if ext == 'tar' else 'container-only')
        aliases = {'gz': ['tgz', 'crate', 'unitypackage'], 'bz2': ['tbz', 'tbz2'],
                   'xz': ['txz'], 'zst': ['tzst'], 'lz': ['tlz'], 'lzo': ['tzo'], 'z': ['taz']}
        for codec in STREAMS:
            def create(codec=codec):
                data = compress(tar, codec)
                for ext in ['tar.' + codec, *aliases.get(codec, [])]:
                    self.add('archives/container.' + ext, data, ext, 'archives',
                             {'kind': 'tar', 'codec': codec, 'children': self.children},
                             scope='native' if ext.startswith('tar.') else 'alias')
            self.optional('tar.' + codec, create)
        for fmt, extensions in [('7z', ['7z', 'cb7']), ('rar', ['rar', 'cbr']),
                                ('wim', ['wim'])]:
            def create(fmt=fmt, extensions=extensions):
                with tempfile.TemporaryDirectory() as temp:
                    temp = Path(temp)
                    for name, data in self.entries.items():
                        target = temp / 'input' / name
                        if name.endswith('/'):
                            target.mkdir(parents=True, exist_ok=True)
                        else:
                            target.parent.mkdir(parents=True, exist_ok=True)
                            target.write_bytes(data)
                    output = temp / ('output.' + fmt)
                    if fmt == 'rar':
                        command(['rar', 'a', '-idq', '-m1', str(output), '.'], cwd=temp / 'input')
                    else:
                        command(['7zz', 'a', '-t' + fmt, '-mx=1', str(output), '.'], cwd=temp / 'input')
                    for ext in extensions:
                        self.add('archives/container.' + ext, output.read_bytes(), ext, 'archives',
                                 {'kind': 'rar' if fmt == 'rar' else '7z', 'children': self.children},
                                 scope='native' if ext == fmt else 'alias')
            self.optional(fmt, create)
        def cab():
            name = b'payload.bin\0'
            raw = self.payload[:30000]
            file_offset = 44
            content_offset = file_offset + 16 + len(name)
            length = content_offset + 8 + len(raw)
            header = struct.pack('<4s5IBB5H', b'MSCF', 0, length, 0, file_offset, 0, 3, 1, 1, 1, 0, 1234, 0)
            folder = struct.pack('<IHH', content_offset, 1, 0)
            file = struct.pack('<II4H', len(raw), 0, 0, 0, 0, 0x20) + name
            block = struct.pack('<IHH', 0, len(raw), len(raw)) + raw
            self.add('archives/container.cab', header + folder + file + block, 'cab', 'archives', {'kind': '7z'})
        self.optional('cab', cab)

    def packages(self):
        import zipfile
        for ext, mime in [('odt', 'text'), ('ods', 'spreadsheet'), ('odp', 'presentation')]:
            mime = 'application/vnd.oasis.opendocument.' + mime
            body = {'odt': '<office:text><text:p>Synthetic text</text:p></office:text>',
                    'ods': '<office:spreadsheet><table:table table:name="Synthetic">'
                           '<table:table-row><table:table-cell office:value-type="string">'
                           '<text:p>Synthetic</text:p></table:table-cell></table:table-row>'
                           '</table:table></office:spreadsheet>',
                    'odp': '<office:presentation><draw:page draw:name="Synthetic"/>'
                           '</office:presentation>'}[ext]
            content = ('<?xml version="1.0"?><office:document-content '
                       'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
                       'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
                       'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
                       'xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0" '
                       'office:version="1.2"><office:body>' + body + '</office:body></office:document-content>').encode()
            entries = {'mimetype': mime.encode(), 'content.xml': content,
                       'META-INF/manifest.xml': ('<manifest:manifest '
                       'xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" '
                       'manifest:version="1.2"><manifest:file-entry manifest:full-path="/" '
                       'manifest:media-type="' + mime + '"/><manifest:file-entry '
                       'manifest:full-path="content.xml" manifest:media-type="text/xml"/>'
                       '</manifest:manifest>').encode()}
            self.add('packages/document.' + ext, zip_bytes(entries, mimetype='mimetype'), ext,
                     'packages', {'kind': 'zip', 'children': {'content.xml': {'kind': 'xml'}},
                                  'stored_first': ['mimetype']})
            if ext == 'odt':
                signed = {**entries, 'META-INF/documentsignatures.xml': b'<document-signatures/>'}
                self.add('controls/signature-marker.odt', zip_bytes(signed, mimetype='mimetype'),
                         'odt', 'packages', {'kind': 'bytes'}, tier='control', expectation='unchanged',
                         note='Valid ODF envelope with a signature marker; rejection guard, not cryptographic signing')
        entries = {'mimetype': b'application/epub+zip', 'META-INF/container.xml':
                   b'<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
                   b'<rootfiles><rootfile full-path="content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>',
                   'content.opf': b'<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="id">'
                   b'<metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:identifier id="id">frbench</dc:identifier>'
                   b'<dc:title>Synthetic</dc:title><dc:language>en</dc:language></metadata>'
                   b'<manifest><item id="chapter" href="chapter.xhtml" media-type="application/xhtml+xml"/></manifest>'
                   b'<spine><itemref idref="chapter"/></spine></package>',
                   'chapter.xhtml': b'<html xmlns="http://www.w3.org/1999/xhtml"><head><title>Synthetic</title></head>'
                   b'<body><p>Synthetic reading sample.</p></body></html>'}
        self.add('packages/book.epub', zip_bytes(entries, mimetype='mimetype'), 'epub', 'packages',
                 {'kind': 'zip', 'children': {'chapter.xhtml': {'kind': 'xml'}}, 'stored_first': ['mimetype']})
        for ext, creator in [('docx', self.docx), ('xlsx', self.xlsx), ('pptx', self.pptx)]:
            def create(ext=ext, creator=creator):
                path = self.directory / ('packages/document.' + ext)
                path.parent.mkdir(parents=True, exist_ok=True)
                creator(path)
                with zipfile.ZipFile(path) as original:
                    items = {i.filename: original.read(i) for i in original.infolist()}
                self.add('packages/document.' + ext, zip_bytes(items), ext, 'packages',
                         {'kind': 'zip', 'children': {k: {'kind': 'xml'} for k in items
                          if k.endswith(('.xml', '.rels'))} | {k: {'kind': 'raster'} for k in items
                          if k.endswith(('.png', '.jpeg', '.jpg'))}}, tier='smoke')
            self.optional(ext, create)
            if ext in ('docx', 'pptx'):
                import importlib.metadata
                package = 'python-docx' if ext == 'docx' else 'python-pptx'
                for case in self.cases:
                    if case['id'] == 'packages-document-' + ext:
                        case['license'] = 'BSD-3-Clause AND MIT'
                        case['provenance']['template'] = {
                            'package': package, 'version': importlib.metadata.version(package),
                            'license_file': 'licenses/template-' + package + '-LICENSE'}

    def docx(self, path):
        from docx import Document
        doc = Document()
        doc.add_heading('Synthetic compression benchmark', 0)
        for i in range(30 * self.scale):
            doc.add_paragraph('Paragraph %d: repetition with Unicode текст and whitespace.' % i)
        doc.save(path)

    def xlsx(self, path):
        from openpyxl import Workbook
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(['id', 'group', 'value', 'formula'])
        for i in range(300 * self.scale):
            sheet.append([i, i % 11, 'repeated string', '=A%d*2' % (i + 2)])
        workbook.save(path)

    def pptx(self, path):
        from pptx import Presentation
        deck = Presentation()
        for i in range(4):
            slide = deck.slides.add_slide(deck.slide_layouts[1])
            slide.shapes.title.text = 'Synthetic slide %d' % i
            slide.placeholders[1].text = 'A repeated text paragraph\n' * 20
        deck.save(path)

    def native(self):
        raw = b'\x12\x34' * (128 * 128 * self.scale)
        header = ('NRRD0005\n# synthetic array\ntype: uint16\ndimension: 2\nsizes: 128 %d\n'
                  'endian: big\nspace origin: (1,2,3)\nlabel:=synthetic\nencoding: raw\n\n' % (128 * self.scale)).encode()
        self.add('native/volume.nrrd', header + raw, 'nrrd', 'scientific', {'kind': 'nrrd'}, tier='smoke')
        self.add('native/weak.nrrd', header.replace(b'encoding: raw', b'encoding: gzip') + gzip.compress(raw, compresslevel=1, mtime=0),
                 'nrrd', 'scientific', {'kind': 'nrrd'})
        blend_header = b'BLENDER-v279'
        dna = b'SDNA' + b'NAME' + struct.pack('<I', 0) + b'TYPE' + struct.pack('<I', 0) + b'TLENSTRC' + struct.pack('<I', 0)
        blend = blend_header + struct.pack('<4siQii', b'DATA', len(self.payload), 4096, 0, 1) + self.payload
        blend += struct.pack('<4siQii', b'DNA1', len(dna), 0, 0, 1) + dna + struct.pack('<4siQii', b'ENDB', 0, 0, 0, 0)
        self.add('native/minimal.blend', blend, 'blend', 'native', {'kind': 'blend'}, scope='syntax',
                 note='Structural Blender stream; not an application-renderable scene')
        psb = struct.pack('>4sH6sHIIHH', b'8BPS', 2, b'\0' * 6, 1, 128, 128, 8, 1)
        psb += struct.pack('>IIQH', 0, 0, 0, 2) + zlib.compress(b'\x2a' * 16384, 1)
        self.add('native/composite.psb', psb, 'psb', 'images', {'kind': 'psb'})
        hdr = bytearray(128)
        struct.pack_into('<I5H', hdr, 0, 0, 0xA5E0, 1, 128, 128, 32)
        layer = struct.pack('<6HB3sH', 1, 0, 0, 0, 0, 0, 255, b'\0' * 3, 5) + b'Layer'
        cel = struct.pack('<HhhBHh5sHH', 0, 0, 0, 255, 2, 0, b'\0' * 5, 128, 128)
        cel += zlib.compress(b'\0\x10\x20\xff' * 16384, 1)
        chunks = b''.join(struct.pack('<IH', len(data) + 6, tag) + data for tag, data
                          in [(0x2004, layer), (0x2005, cel), (0x7777, b'unknown metadata')])
        frame = struct.pack('<IHHHHI', len(chunks) + 16, 0xF1FA, 3, 123, 0, 0) + chunks
        struct.pack_into('<I', hdr, 0, 128 + len(frame))
        for ext in ['ase', 'aseprite']:
            self.add('native/frame.' + ext, bytes(hdr) + frame, ext, 'images', {'kind': 'aseprite'})
        body = b'\x08\0\0\x18\x01\0'
        meta = b'<metadata>Synthetic frame metadata</metadata>' * 1000
        body += struct.pack('<HI', (77 << 6) | 63, len(meta)) + meta + b'\0\0'
        swf = b'CWS\x0a' + struct.pack('<I', len(body) + 8) + zlib.compress(body, 1)
        self.add('native/animation.swf', swf, 'swf', 'native', {'kind': 'swf'})
        tgs = json.dumps({'tgs': 1, 'v': '5.5.2', 'w': 512, 'h': 512, 'fr': 60,
                          'ip': 0, 'op': 180, 'layers': [], 'assets': [],
                          'extra': ['Synthetic metadata'] * 300}, indent=4).encode()
        self.add('native/sticker.tgs', gzip.compress(tgs, compresslevel=1, mtime=0), 'tgs', 'native', {'kind': 'tgs'})
        for ext, data in [('psb', psb), ('aseprite', bytes(hdr) + frame), ('swf', swf),
                          ('nrrd', header + raw), ('tgs', gzip.compress(tgs, mtime=0))]:
            self.add('controls/truncated.' + ext, data[:len(data) // 2], ext, 'native',
                     {'kind': 'bytes'}, expectation='unchanged', tier='control')

    def sqlite(self):
        for ext in ['sqlite', 'sqlite3', 'db', 'qgd', 'vscdb', 'sqlitedb', 'gpkg', 'mbtiles']:
            path = self.directory / ('data/slack.' + ext)
            path.parent.mkdir(parents=True, exist_ok=True)
            with sqlite3.connect(path) as db:
                db.execute('PRAGMA application_id=123')
                db.execute('PRAGMA user_version=7')
                db.execute('CREATE TABLE measurements(id INTEGER PRIMARY KEY,value TEXT)')
                db.executemany('INSERT INTO measurements VALUES (?,?)',
                               [(i, 'stored synthetic data ' * 20) for i in range(2000 * self.scale)])
                db.execute('CREATE INDEX values_idx ON measurements(value)')
                db.execute('DELETE FROM measurements WHERE id >= 50')
            self.add('data/slack.' + ext, None, ext, 'data', {'kind': 'sqlite'},
                     scope='container-only' if ext in ('gpkg', 'mbtiles') else 'native',
                     tier='smoke' if ext == 'sqlite' else 'extended')

    def columnar(self):
        import pyarrow as pa
        import pyarrow.feather as feather
        import pyarrow.orc as orc
        import pyarrow.parquet as pq
        from decimal import Decimal
        table = pa.table({'id': list(range(10000 * self.scale)),
                          'label': ['repeat label' if i % 13 else None for i in range(10000 * self.scale)],
                          'precise': pa.array([Decimal('1.2300')] * (10000 * self.scale), type=pa.decimal128(12, 4))})
        table = table.replace_schema_metadata({b'corpus': b'synthetic', b'contract': b'preserve'})
        for ext in ['parquet', 'orc', 'feather', 'arrow', 'ipc']:
            path = self.directory / ('data/uncompressed.' + ext)
            path.parent.mkdir(parents=True, exist_ok=True)
            # ORC cannot represent arbitrary Arrow schema metadata.
            data = table.replace_schema_metadata(None) if ext == 'orc' else table
            if ext == 'parquet':
                pq.write_table(data, path, compression='NONE', use_dictionary=False, row_group_size=1000)
            elif ext == 'orc':
                orc.write_table(data, path, compression='uncompressed')
            elif ext == 'feather':
                feather.write_feather(data, path, compression='uncompressed')
            else:
                with pa.OSFile(str(path), 'wb') as sink:
                    with pa.ipc.new_file(sink, data.schema) as writer:
                        writer.write_table(data)
            self.add('data/uncompressed.' + ext, None, ext, 'data',
                     {'kind': 'arrow', 'format': ext}, tier='smoke' if ext == 'parquet' else 'extended')
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / 'compact.parquet'
            pq.write_table(table, path, compression='zstd', compression_level=19)
            self.add('data/compact.parquet', path.read_bytes(), 'parquet', 'data', {'kind': 'arrow', 'format': 'parquet'})

    def avro(self):
        import fastavro
        schema = {'type': 'record', 'name': 'SyntheticMeasurement', 'fields': [
            {'name': 'id', 'type': 'long'}, {'name': 'label', 'type': 'string'}]}
        buffer = io.BytesIO()
        fastavro.writer(buffer, schema, ({'id': i, 'label': 'repeat label'}
                                        for i in range(5000 * self.scale)), codec='null',
                        sync_marker=b'frbench-fixed-01')
        self.add('data/uncompressed.avro', buffer.getvalue(), 'avro', 'data', {'kind': 'avro'})

    def scientific(self):
        import numpy as np
        import h5py
        import netCDF4
        import tifffile
        from astropy.io import fits
        data = np.zeros((256, 256), dtype='int16')
        data[0] = np.arange(256)
        for ext in ['h5', 'hdf', 'hdf5']:
            path = self.directory / ('scientific/arrays.' + ext)
            path.parent.mkdir(parents=True, exist_ok=True)
            with h5py.File(path, 'w') as file:
                file.attrs['label'] = 'synthetic metadata'
                values = file.create_dataset('values', data=data, chunks=(32, 32))
                values.attrs['unit'] = 'arbitrary'
            self.add('scientific/arrays.' + ext, None, ext, 'scientific', {'kind': 'hdf5'})
        for ext in ['nc', 'nc4']:
            path = self.directory / ('scientific/arrays.' + ext)
            with netCDF4.Dataset(path, 'w', format='NETCDF4') as file:
                file.createDimension('x', 256)
                file.createDimension('y', 256)
                file.title = 'Synthetic metadata'
                values = file.createVariable('values', 'i2', ('x', 'y'))
                values.units = 'arbitrary'
                values[:] = data
            self.add('scientific/arrays.' + ext, None, ext, 'scientific', {'kind': 'netcdf'})
        for ext in ['tif', 'tiff']:
            path = self.directory / ('scientific/multipage.' + ext)
            with tifffile.TiffWriter(path) as file:
                file.write(data, metadata=None, description='Synthetic image')
                file.write(data * 2, metadata=None, description='Second page')
            self.add('scientific/multipage.' + ext, None, ext, 'scientific', {'kind': 'tiff'})
        for ext in ['fits', 'fit', 'fts']:
            path = self.directory / ('scientific/arrays.' + ext)
            hdu = fits.PrimaryHDU(data)
            hdu.header['OBJECT'] = ('synthetic array', 'preserve comment')
            hdu.header.add_history('Synthetic creation')
            hdu.writeto(path, overwrite=True, checksum=True)
            self.add('scientific/arrays.' + ext, None, ext, 'scientific', {'kind': 'fits'})

    def serialization(self):
        import numpy as np
        import scipy.io
        import torch
        import pandas as pd
        import pyreadstat
        import zarr
        from numcodecs import Zlib
        path = self.directory / 'scientific/arrays.mat'
        path.parent.mkdir(parents=True, exist_ok=True)
        scipy.io.savemat(path, {'values': np.zeros((256, 256)), 'label': 'synthetic'}, do_compression=True)
        self.add('scientific/arrays.mat', None, 'mat', 'scientific', {'kind': 'mat'})
        for ext in ['pt', 'pth']:
            path = self.directory / ('scientific/state.' + ext)
            base = torch.zeros(5000)
            torch.save({**{f'parameter_{i:05}': base for i in range(1500)}, 'view': base[3::2]}, path)
            self.add('scientific/state.' + ext, None, ext, 'scientific', {'kind': 'checkpoint'})
        for ext, compression in [('sav', False), ('zsav', True)]:
            path = self.directory / ('scientific/records.' + ext)
            frame = pd.DataFrame({'value': [1.0] * 5000, 'label': ['synthetic'] * 5000})
            pyreadstat.write_sav(frame, path, compress=compression, row_compress=False,
                                file_label='Synthetic measurements', column_labels=['Value', 'Label'])
            self.add('scientific/records.' + ext, None, ext, 'scientific', {'kind': 'spss'})
        path = self.directory / 'scientific/arrays.zarr'
        group = zarr.open_group(str(path), mode='w')
        group.attrs['label'] = 'synthetic'
        group.create_dataset('values', data=np.zeros((256, 256), dtype='int16'),
                             chunks=(32, 32), compressor=Zlib(level=1))
        self.add('scientific/arrays.zarr', None, 'zarr', 'scientific', {'kind': 'zarr'})
        # Valid NPZ is a real NPY ZIP corpus, unlike generic archive aliases.
        buffer = io.BytesIO()
        np.savez(buffer, values=np.zeros((128, 128)), labels=np.arange(128))
        self.add('packages/arrays.npz', buffer.getvalue(), 'npz', 'packages', {'kind': 'npz'})

    def r(self):
        for ext in ['rds', 'rda', 'rdata']:
            path = self.directory / ('scientific/arrays.' + ext)
            path.parent.mkdir(parents=True, exist_ok=True)
            expression = ('x<-list(i=rep(1:100,1000), f=c(0,-0,NaN,NA_real_,Inf), label="synthetic");')
            expression += (f'saveRDS(x,file={json.dumps(str(path))},compress=FALSE,version=2)' if ext == 'rds'
                           else f'save(x,file={json.dumps(str(path))},compress=FALSE,version=2)')
            command(['Rscript', '-e', expression])
            self.add('scientific/arrays.' + ext, None, ext, 'scientific', {'kind': 'r'})
            raw = path.read_bytes()
            self.add('scientific/weak.' + ext, gzip.compress(raw, compresslevel=1, mtime=0), ext,
                     'scientific', {'kind': 'r'})

    def duckdb(self):
        import duckdb
        path = self.directory / 'data/slack.duckdb'
        path.parent.mkdir(parents=True, exist_ok=True)
        with duckdb.connect(str(path)) as db:
            db.execute("CREATE TABLE measurements AS SELECT range AS id, repeat('synthetic',100) AS label FROM range(20000)")
            db.execute('CHECKPOINT')
            db.execute('DELETE FROM measurements WHERE id >= 100')
            db.execute('CHECKPOINT')
        self.add('data/slack.duckdb', None, 'duckdb', 'data', {'kind': 'duckdb'})

    def images(self):
        from PIL import Image, PngImagePlugin
        image = Image.new('RGB', (256, 256))
        image.putdata([(x // 16 * 16, y // 16 * 16, (x + y) // 32 * 16)
                       for y in range(256) for x in range(256)])
        formats = {'png': ('PNG', {'compress_level': 0}), 'jpg': ('JPEG', {'quality': 95}),
                   'gif': ('GIF', {}), 'webp': ('WEBP', {'lossless': True, 'quality': 0}),
                   'bmp': ('BMP', {}), 'dib': ('DIB', {}), 'tga': ('TGA', {}),
                   'ppm': ('PPM', {}), 'pcx': ('PCX', {}), 'ico': ('ICO', {}),
                   'icns': ('ICNS', {}), 'jp2': ('JPEG2000', {'irreversible': False}),
                   'j2k': ('JPEG2000', {'no_jp2': True, 'irreversible': False}),
                   'avif': ('AVIF', {'quality': 100})}
        for ext, (form, options) in formats.items():
            def create(ext=ext, form=form, options=options):
                buffer = io.BytesIO()
                image.save(buffer, form, **options)
                self.add('images/blocks.' + ext, buffer.getvalue(), ext, 'images', {'kind': 'raster'},
                         tier='smoke' if ext == 'png' else 'extended')
            self.optional(ext, create)
        buffer = io.BytesIO()
        image.save(buffer, 'PNG', compress_level=9)
        self.add('images/compact.png', buffer.getvalue(), 'png', 'images', {'kind': 'raster'})
        info = PngImagePlugin.PngInfo()
        info.add_text('Comment', 'Synthetic metadata retained under keep-meta')
        buffer = io.BytesIO()
        image.save(buffer, 'PNG', pnginfo=info, compress_level=1)
        self.add('images/metadata.png', buffer.getvalue(), 'png', 'images', {'kind': 'raster'})
        for ext in ['gif', 'apng']:
            buffer = io.BytesIO()
            image.save(buffer, 'PNG' if ext == 'apng' else 'GIF', save_all=True,
                       append_images=[image.transpose(Image.Transpose.FLIP_LEFT_RIGHT)],
                       duration=[100, 250], loop=2)
            self.add('images/animated.' + ext, buffer.getvalue(), ext, 'images', {'kind': 'raster'})
        for ext, img in [('pgm', image.convert('L')), ('pbm', image.convert('1')), ('pnm', image)]:
            buffer = io.BytesIO()
            img.save(buffer, 'PPM')
            self.add('images/blocks.' + ext, buffer.getvalue(), ext, 'images', {'kind': 'raster'})
        svg = (b'<svg  xmlns = "http://www.w3.org/2000/svg" width="256" height="256">'
               b'<rect  width = "256" height = "256" fill = "red" />' * 1 + b'</svg>')
        self.add('images/drawing.svg', svg, 'svg', 'images', {'kind': 'svg-render'})
        self.add('images/drawing.svgz', gzip.compress(svg, compresslevel=1, mtime=0), 'svgz', 'images',
                 {'kind': 'svg-render'})
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / 'input.png'
            image.save(source)
            for ext in ['psd', 'exr', 'heic', 'heif', 'jxl']:
                def create(ext=ext):
                    output = Path(temp) / ('output.' + ext)
                    command(['magick', str(source), str(output)])
                    self.add('images/blocks.' + ext, output.read_bytes(), ext, 'images',
                             {'kind': 'imagemagick'} if ext in ('exr', 'jxl', 'psd') else {'kind': 'raster'})
                self.optional(ext, create)

    def dicom(self):
        import numpy as np
        from pydicom.dataset import FileDataset, FileMetaDataset
        from pydicom.uid import ExplicitVRLittleEndian
        meta = FileMetaDataset()
        meta.TransferSyntaxUID = ExplicitVRLittleEndian
        meta.MediaStorageSOPClassUID = '1.2.840.10008.5.1.4.1.1.7'
        meta.MediaStorageSOPInstanceUID = '1.2.826.0.1.3680043.10.999.1'
        data = FileDataset('synthetic', {}, file_meta=meta, preamble=b'\0' * 128)
        data.SOPClassUID = meta.MediaStorageSOPClassUID
        data.SOPInstanceUID = meta.MediaStorageSOPInstanceUID
        data.PatientName, data.PatientID = 'SYNTHETIC^BENCHMARK', 'NOT-A-PATIENT'
        data.Modality = 'OT'
        data.Rows = data.Columns = 128
        data.SamplesPerPixel, data.PhotometricInterpretation = 1, 'MONOCHROME2'
        data.BitsAllocated, data.BitsStored, data.HighBit, data.PixelRepresentation = 16, 16, 15, 0
        data.PixelData = np.zeros((128, 128), dtype='<u2').tobytes()
        buffer = io.BytesIO()
        data.save_as(buffer, enforce_file_format=True)
        for ext in ['dcm', 'dicom', 'dic']:
            self.add('medical/synthetic.' + ext, buffer.getvalue(), ext, 'medical', {'kind': 'dicom'})

    def pdf(self):
        import pikepdf
        pdf = pikepdf.Pdf.new()
        for _ in range(3):
            page = pdf.add_blank_page(page_size=(300, 300))
            content = b'q 1 0 0 rg 20 20 100 100 re f Q\n' * 200
            page.Contents = pdf.make_stream(content)
        pdf.docinfo['/Title'] = 'Synthetic compression benchmark'
        pdf.attachments['readme.txt'] = b'Synthetic attachment preserved byte for byte'
        buffer = io.BytesIO()
        pdf.save(buffer, compress_streams=False)
        for ext in ['pdf', 'ai']:
            self.add('documents/vector.' + ext, buffer.getvalue(), ext, 'documents', {'kind': 'pdf'},
                     scope='native' if ext == 'pdf' else 'alias', tier='smoke' if ext == 'pdf' else 'extended')
        buffer = io.BytesIO()
        pdf.save(buffer, encryption=pikepdf.Encryption(owner='synthetic', user='synthetic', R=6))
        self.add('controls/encrypted.pdf', buffer.getvalue(), 'pdf', 'documents', {'kind': 'bytes'},
                 tier='control', expectation='unchanged')

    def fonts(self):
        from fontTools.fontBuilder import FontBuilder
        from fontTools.pens.ttGlyphPen import TTGlyphPen
        font = FontBuilder(1000, isTTF=True)
        names = ['.notdef', *[f'g{i}' for i in range(128)]]
        font.setupGlyphOrder(names)
        font.setupCharacterMap({i + 32: f'g{i}' for i in range(128)})
        glyphs = {}
        for name in names:
            pen = TTGlyphPen(None)
            pen.moveTo((100, 0))
            pen.lineTo((100, 700))
            pen.lineTo((500, 700))
            pen.lineTo((500, 0))
            pen.closePath()
            glyphs[name] = pen.glyph()
        font.setupGlyf(glyphs)
        font.setupHorizontalMetrics({name: (600, 0) for name in names})
        font.setupHorizontalHeader(ascent=800, descent=-200)
        font.setupNameTable({'familyName': 'Frbench Synthetic', 'styleName': 'Regular',
                            'uniqueFontIdentifier': 'Frbench Synthetic 1',
                            'fullName': 'Frbench Synthetic Regular', 'psName': 'FrbenchSynthetic'})
        font.setupOS2(sTypoAscender=800, sTypoDescender=-200, usWinAscent=800, usWinDescent=200)
        font.setupPost()
        font.font['head'].created = font.font['head'].modified = 2082844800 + 1577934246
        for ext in ['woff', 'woff2']:
            font.font.flavor = ext
            buffer = io.BytesIO()
            font.font.save(buffer)
            self.add('fonts/synthetic.' + ext, buffer.getvalue(), ext, 'fonts', {'kind': 'font'})

    def media(self):
        import math
        import wave
        with tempfile.TemporaryDirectory() as temp:
            temp = Path(temp)
            wav = temp / 'input.wav'
            with wave.open(str(wav), 'wb') as file:
                file.setparams((1, 2, 16000, 0, 'NONE', 'not compressed'))
                file.writeframes(b''.join(struct.pack('<h', int(12000 * math.sin(2 * math.pi * 440 * i / 16000)))
                                          for i in range(32000)))
            audio = {'flac': ['-c:a', 'flac', '-compression_level', '0'], 'm4a': ['-c:a', 'alac'],
                     'm4b': ['-c:a', 'alac', '-f', 'ipod'], 'tta': ['-c:a', 'tta'],
                     'wv': ['-c:a', 'wavpack'], 'ogg': ['-c:a', 'libvorbis', '-q:a', '8'],
                     'oga': ['-c:a', 'flac', '-f', 'ogg'], 'opus': ['-c:a', 'libopus'],
                     'mp3': ['-c:a', 'libmp3lame', '-b:a', '320k']}
            for ext, options in audio.items():
                def create(ext=ext, options=options):
                    output = temp / ('output.' + ext)
                    command(['ffmpeg', '-v', 'error', '-i', str(wav), *options, '-y', str(output)])
                    self.add('audio/tone.' + ext, output.read_bytes(), ext, 'audio', {'kind': 'audio'})
                self.optional(ext, create)
            def ape():
                output = temp / 'output.ape'
                command(['mac', str(wav), str(output), '-c1000'])
                self.add('audio/tone.ape', output.read_bytes(), 'ape', 'audio', {'kind': 'audio'})
            self.optional('ape', ape)
            for ext in ['mp4', 'mkv', 'webm', 'mov', 'm4v', 'avi', 'wmv', 'asf', '3gp', 'ts', 'mts', 'm2ts']:
                def create(ext=ext):
                    output = temp / ('output.' + ext)
                    options = (['-c:v', 'libvpx-vp9', '-lossless', '1', '-deadline', 'realtime'] if ext == 'webm'
                               else ['-c:v', 'ffv1'] if ext == 'avi'
                               else ['-c:v', 'wmv2'] if ext in ('wmv', 'asf')
                               else ['-c:v', 'libx264', '-crf', '0', '-preset', 'ultrafast'])
                    if ext in ('ts', 'mts', 'm2ts'):
                        options += ['-f', 'mpegts']
                    command(['ffmpeg', '-v', 'error', '-f', 'lavfi', '-i',
                             'testsrc2=size=128x96:rate=12:duration=1', *options, '-pix_fmt', 'yuv420p', '-an', '-y', str(output)])
                    self.add('video/test-pattern.' + ext, output.read_bytes(), ext, 'video', {'kind': 'video'})
                self.optional(ext, create)

    def extra_images(self):
        from PIL import Image
        import numpy as np
        import tifffile
        image = Image.new('RGB', (32, 32), (12, 34, 56))
        buffer = io.BytesIO()
        image.save(buffer, 'ICO', sizes=[(32, 32)], bitmap_format='bmp')
        cur = bytearray(buffer.getvalue())
        struct.pack_into('<H', cur, 2, 2)
        struct.pack_into('<HH', cur, 10, 4, 5)
        self.add('images/cursor.cur', bytes(cur), 'cur', 'images', {'kind': 'raster'})
        pages = []
        for frame in [image, Image.new('RGB', (32, 32), (56, 34, 12))]:
            buffer = io.BytesIO()
            frame.save(buffer, 'PCX')
            pages.append(buffer.getvalue())
        offset = 4100
        offsets = []
        for page in pages:
            offsets.append(offset)
            offset += len(page)
        dcx = struct.pack('<I', 0x3ADE68B1) + struct.pack('<1024I', *(offsets + [0] * (1024 - len(offsets)))) + b''.join(pages)
        self.add('images/multipage.dcx', dcx, 'dcx', 'images', {'kind': 'raster'})
        path = self.directory / 'images/raw.dng'
        tifffile.imwrite(path, np.zeros((128, 128), dtype='uint16'), photometric=32803, metadata=None,
                          extratags=[(50706, 'B', 4, (1, 4, 0, 0), False),
                                     (50707, 'B', 4, (1, 1, 0, 0), False),
                                     (50708, 's', 0, 'Frbench Synthetic Camera', False),
                                     (33421, 'H', 2, (2, 2), False),
                                     (33422, 'B', 4, (0, 1, 1, 2), False)])
        self.add('images/raw.dng', None, 'dng', 'images', {'kind': 'tiff'},
                 note='Minimal synthetic DNG mosaic; generated camera metadata')

    def new_native(self):
        if 'warc' in self.catalog['supported_extensions']:
            records = []
            for i in range(3):
                raw = self.payload if i < 2 else b''
                record = (b'WARC/1.1\r\nWARC-Type: resource\r\nWARC-Record-ID: '
                          + ('<urn:uuid:00000000-0000-0000-0000-%012d>' % i).encode()
                          + b'\r\nWARC-Date: 2020-01-02T03:04:06Z\r\n'
                          b'WARC-Target-URI: https://example.org/synthetic\r\nContent-Length: '
                          + str(len(raw)).encode() + b'\r\n\r\n' + raw + b'\r\n\r\n')
                records.append(record)
            self.add('native/capture.warc', b''.join(records), 'warc', 'native', {'kind': 'warc'})
            self.add('native/weak.warc.gz', b''.join(gzip.compress(r, compresslevel=1, mtime=0) for r in records),
                     'warc', 'native', {'kind': 'warc'})
        if 'nib' in self.catalog['supported_extensions']:
            def integer(value):
                result = bytearray()
                while value >= 128:
                    result.append(value & 127)
                    value >>= 7
                result.append(value | 128)
                return bytes(result)
            keys = integer(5) + b'value'
            value = integer(0) + b'\x08' + integer(300) + b'synthetic ' * 30
            values = value * 20
            objects = b''.join(integer(0) + integer(i) + integer(1) for i in range(20))
            classes = integer(9) + integer(0) + b'NSObject\0'
            key_offset = 50 + len(objects)
            value_offset = key_offset + len(keys)
            class_offset = value_offset + len(values)
            nib = b'NIBArchive' + struct.pack('<10I', 1, 9, 20, 50, 1, key_offset, 20, value_offset, 1, class_offset)
            self.add('native/duplicate-values.nib', nib + objects + keys + values + classes,
                     'nib', 'native', {'kind': 'nib'}, scope='syntax')
        # USDZ stores the first scene and aligns every member to a 64-byte boundary.
        import zipfile
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, 'w') as archive:
            name = 'scene.usda'
            info = zipfile.ZipInfo(name, FIXED_TIME)
            padding = (-(30 + len(name) + 4)) % 64
            info.extra = struct.pack('<HH', 0xCAFE, padding) + b'\0' * padding
            archive.writestr(info, b'#usda 1.0\ndef Xform "Synthetic" {}\n')
        self.add('packages/scene.usdz', buffer.getvalue(), 'usdz', 'packages', {'kind': 'zip'})

    def additional_formats(self):
        """Add native profiles for registry routes and compound filename dispatch."""
        payload = self.payload
        self.add('native/catalog.car', car_profile(payload), 'car', 'native', {'kind': 'car'},
                 note='Synthetic BOMStore catalog with a compressed MLEC rendition')
        cpio = cpio_profile(payload, self.json, self.xml)
        self.add('archives/cpio-tree.cpio', cpio, 'cpio', 'archives', {'kind': 'cpio'},
                 note='newc CPIO tree with regular files, a directory, symlink and hard links')
        compressed_cpio = cpio_bz2(payload, self.json, self.xml)
        self.add('streams/cpio-tree.cpbz2', compressed_cpio, 'cpbz2', 'archives',
                 {'kind': 'cpbz2'}, note='newc CPIO tree wrapped in a weak bzip2 stream')
        self.add('streams/cpio-tree.cpio.bz2', compressed_cpio, 'cpio.bz2', 'archives',
                 {'kind': 'cpbz2'}, scope='alias',
                 provenance={'type': 'derived-alias', 'parent': 'streams-cpio-tree-cpbz2',
                             'description': 'Same CPIO+bzip2 bitstream under the compound suffix'})

        self.odf_variants()

        for ext in ['rds', 'rda', 'rdata']:
            source = self.directory / f'scientific/arrays.{ext}'
            if not source.is_file():
                continue
            raw = source.read_bytes()
            for codec in ['gz', 'bz2', 'xz']:
                self.add(f'scientific/arrays.{ext}.{codec}', compress(raw, codec, 1), ext,
                         'scientific', {'kind': 'r'},
                         note=f'R serialization in its supported {codec} compound suffix')

        crx_entries = {**self.entries, 'manifest.json': json.dumps({
            'manifest_version': 3, 'name': 'Frbench synthetic extension', 'version': '1.0.0',
            'description': 'Generated data for a compression benchmark.'}, indent=2).encode()}
        crx_children = {**self.children, 'manifest.json': {'kind': 'json'}}
        self.optional('crx', lambda: self.add(
            'archives/signed-extension.crx', crx3(zip_bytes(crx_entries, level=1)),
            'crx', 'archives', {'kind': 'crx', 'children': crx_children},
            note='CRX3 with a temporary test-only signing key and a minimal extension manifest'))

    def odf_variants(self):
        """Build standards-shaped ODF package profiles for the remaining ODF suffixes."""
        namespace = {
            'office': 'urn:oasis:names:tc:opendocument:xmlns:office:1.0',
            'text': 'urn:oasis:names:tc:opendocument:xmlns:text:1.0',
            'table': 'urn:oasis:names:tc:opendocument:xmlns:table:1.0',
            'draw': 'urn:oasis:names:tc:opendocument:xmlns:drawing:1.0',
            'chart': 'urn:oasis:names:tc:opendocument:xmlns:chart:1.0',
            'math': 'http://www.w3.org/1998/Math/MathML',
        }
        profiles = {
            'odg': ('graphics', '<office:drawing><draw:page draw:name="Synthetic"/></office:drawing>'),
            'otg': ('graphics-template', '<office:drawing><draw:page draw:name="Template"/></office:drawing>'),
            'odf': ('formula', '<office:formula><math:math><math:mi>x</math:mi></math:math></office:formula>'),
            'odb': ('database', '<office:database/>'),
            'odc': ('chart', '<office:chart><chart:chart/></office:chart>'),
            'odi': ('image', '<office:image/>'),
            'odm': ('text-master', '<office:text><text:p>Synthetic master text</text:p></office:text>'),
            'ott': ('text-template', '<office:text><text:p>Synthetic text template</text:p></office:text>'),
            'ots': ('spreadsheet-template', '<office:spreadsheet><table:table table:name="Template"/></office:spreadsheet>'),
            'otp': ('presentation-template', '<office:presentation><draw:page draw:name="Template"/></office:presentation>'),
            'oth': ('text-web', '<office:text><text:p>Synthetic web template</text:p></office:text>'),
            'otm': ('text-master-template', '<office:text><text:p>Synthetic master template</text:p></office:text>'),
            'otc': ('chart-template', '<office:chart><chart:chart/></office:chart>'),
            'oti': ('image-template', '<office:image/>'),
            'otf': ('formula-template', '<office:formula><math:math><math:mi>x</math:mi></math:math></office:formula>'),
        }
        for ext, (document_type, body) in profiles.items():
            mime = 'application/vnd.oasis.opendocument.' + document_type
            content = ('<?xml version="1.0" encoding="UTF-8"?>'
                       '<office:document-content ' + ' '.join(
                           f'xmlns:{prefix}="{uri}"' for prefix, uri in namespace.items()) +
                       ' office:version="1.2"><office:body>' + body +
                       '</office:body></office:document-content>').encode()
            manifest = ('<manifest:manifest '
                        'xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" '
                        'manifest:version="1.2"><manifest:file-entry manifest:full-path="/" '
                        'manifest:media-type="' + mime + '"/><manifest:file-entry '
                        'manifest:full-path="content.xml" manifest:media-type="text/xml"/>'
                        '</manifest:manifest>').encode()
            archive = {'mimetype': mime.encode(), 'content.xml': content,
                       'META-INF/manifest.xml': manifest}
            self.add('packages/native.' + ext, zip_bytes(archive, mimetype='mimetype'), ext,
                     'packages', {'kind': 'zip', 'children': {'content.xml': {'kind': 'xml'}},
                                  'stored_first': ['mimetype']},
                     note=f'ODF {document_type} package with its registered media type')

    def aliases(self):
        groups = {
            'jpg': ['jpeg', 'jpe', 'jfif', 'jif', 'jfi', 'thm'],
            'jp2': ['jpf', 'jpx'], 'tga': ['targa'], 'ico': [],
            'doc': ['dot'], 'xls': ['xlt', 'xla'], 'ppt': ['pot', 'pps']}
        for source_ext, aliases in groups.items():
            samples = [c for c in self.cases if c['extension'] == source_ext and c['tier'] != 'control']
            if samples:
                source = samples[0]
                for ext in aliases:
                    self.add('aliases/sample.' + ext, (self.root / source['path']).read_bytes(), ext,
                             source['family'], source['oracle'], scope='alias',
                             provenance={'type': 'derived-alias', 'parent': source['id'],
                                         'description': 'Same bitstream with a routing alias; no app-profile claim'},
                             license=source['license'])


def handler_catalog(source, constants):
    formats = ast.parse((source / 'filerepack/formats.py').read_text())
    aliases = {}
    for node in formats.body:
        if isinstance(node, ast.Assign) and any(isinstance(t, ast.Name) and t.id == 'STANDALONE_ALIASES' for t in node.targets):
            aliases = ast.literal_eval(node.value)
    dispatch = ast.parse((source / 'filerepack/dispatch.py').read_text())
    packers = {}
    for node in dispatch.body:
        if isinstance(node, ast.AnnAssign) and isinstance(node.target, ast.Name) and node.target.id == '_PACKERS':
            packers = {ast.literal_eval(k): v.args[0].id for k, v in zip(node.value.keys, node.value.values)}
    from .catalog import SPECIAL_FAMILY
    result = {ext: 'archive:' + SPECIAL_FAMILY.get(ext, 'zip') for ext in constants['ARCHIVE_EXTS']}
    result.update({ext: packers.get(aliases.get(ext, ext), 'unmapped:' + ext) for ext in constants['STANDALONE_EXTS']})
    return result


def catalog_from_source(path):
    source = Path(path)
    constants = ast.parse((source / 'filerepack/consts.py').read_text())
    values = {}
    for node in constants.body:
        if isinstance(node, ast.Assign):
            for target in node.targets:
                if isinstance(target, ast.Name) and target.id in ('ARCHIVE_EXTS', 'STANDALONE_EXTS'):
                    values[target.id] = ast.literal_eval(node.value)
    return {'archive_extensions': values['ARCHIVE_EXTS'], 'standalone_extensions': values['STANDALONE_EXTS'],
            'supported_extensions': sorted(set(values['ARCHIVE_EXTS'] + values['STANDALONE_EXTS'])),
            'handler_by_extension': handler_catalog(source, values)}


def generate(args):
    root = Path(args.root).resolve()
    catalog_path = root / 'corpus/catalog.json'
    if args.filerepack:
        catalog = catalog_from_source(args.filerepack)
        write_json(catalog_path, catalog)
    else:
        catalog = json.loads(catalog_path.read_text())
    if getattr(args, 'extend', False):
        return extend(root, catalog, args)
    # Owned generated tree can be recreated; original pinned assets are separate.
    directory = root / 'corpus/generated'
    if directory.exists():
        shutil.rmtree(directory)
    builder = Builder(root, catalog, args.scale)
    qgs = builder.text()
    builder.streams()
    builder.archives(qgs)
    builder.native()
    builder.sqlite()
    if not args.minimal:
        for name in ['packages', 'columnar', 'avro', 'scientific', 'serialization', 'r',
                     'duckdb', 'images', 'extra_images', 'new_native', 'dicom', 'pdf', 'fonts', 'media']:
            builder.optional(name, getattr(builder, name))
    else:
        builder.gaps.append({'generator': 'optional', 'reason': 'Skipped by --minimal'})
    builder.additional_formats()
    originals_path = root / 'corpus/originals/index.json'
    if originals_path.exists():
        builder.cases.extend(json.loads(originals_path.read_text())['cases'])
    builder.aliases()
    manifest = {'schema_version': 1, 'seed': 1729, 'scale': args.scale,
                'supported_extensions': catalog['supported_extensions'],
                'cases': sorted(builder.cases, key=lambda c: c['id']), 'generation_gaps': builder.gaps,
                'generated_license': 'BSD-3-Clause'}
    write_json(root / 'corpus/manifest.json', manifest)
    # Remove any partial optional files not indexed in manifest.
    referenced = {(root / c['path']).resolve() for c in builder.cases}
    for path in sorted(directory.rglob('*'), reverse=True):
        if path.is_file() and not any(path == p or p in path.parents for p in referenced):
            path.unlink()
    print(f"Generated {len(builder.cases)} cases; {len(builder.gaps)} explicit generation gaps.")
    return 1 if args.strict and builder.gaps else 0


def extend(root, catalog, args):
    """Add newly generated coverage while preserving every existing pinned specimen."""
    from .common import read_manifest

    root = Path(root)
    builder = Builder(root, catalog, args.scale)
    builder.additional_formats()
    previous = read_manifest(root)
    by_id = {case['id']: case for case in previous['cases']}
    by_id.update({case['id']: case for case in builder.cases})
    gaps = {item['generator']: item for item in previous.get('generation_gaps', [])}
    gaps.update({item['generator']: item for item in builder.gaps})
    manifest = {**previous, 'supported_extensions': catalog['supported_extensions'],
                'cases': sorted(by_id.values(), key=lambda case: case['id']),
                'generation_gaps': [gaps[name] for name in sorted(gaps)]}
    write_json(root / 'corpus/manifest.json', manifest)
    print(f"Extended corpus to {len(manifest['cases'])} cases; {len(builder.cases)} additions/updates; "
          f"{len(builder.gaps)} new generation gaps.")
    return 1 if args.strict and builder.gaps else 0
